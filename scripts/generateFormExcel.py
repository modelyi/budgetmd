#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
生成脚本（第二步）：读取 scripts/generateFormExcel.ts 输出的中间 JSON
（docs/formexcel/.dump/<表单编码>.json），用 openpyxl 统一写出带样式 xlsx
（docs/formexcel/<表单编码>.xlsx）。

样式约定：
  - #002f6c / #001e4a 表头：深蓝填充 + 白色加粗字体 + 水平垂直居中；
  - #475569（只读列）：灰蓝填充 + 白色加粗；
  - #b45309（可修改列）：琥珀填充 + 白色加粗；
  - 普通数据格：不填色、黑字、顶端对齐；
  - 全表细边框（thin, #e2e8f0）。

用法：python scripts/generateFormExcel.py
"""
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.normpath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)), '..', 'docs', 'formexcel'))
DUMP_DIR = os.path.join(BASE_DIR, '.dump')

HEADER_FILLS = {
    '#002f6c': 'FF002F6C',  # 深蓝主表头
    '#001e4a': 'FF001E4A',  # 深蓝合计表头
    '#475569': 'FF475569',  # 只读列灰蓝
    '#b45309': 'FFB45309',  # 可修改列琥珀
}

THIN_SIDE = Side(style='thin', color='E2E8F0')
ALL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFFFF')
DATA_FONT = Font(name='Calibri', size=11, color='FF000000')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
DATA_ALIGN = Alignment(vertical='top')

ILLEGAL_SHEET_NAME_CHARS = re.compile(r'[\\/?*\[\]:]')


def sanitize_sheet_name(name, fallback):
    """sheet 名清洗：非法字符（\\ / ? * [ ] :）替换为空格，限长 31。"""
    s = ILLEGAL_SHEET_NAME_CHARS.sub(' ', str(name or fallback))
    s = re.sub(r'\s+', ' ', s).strip()
    if not s:
        s = str(fallback)
    if len(s) > 31:
        s = s[:31].strip()
    return s


def unique_sheet_name(name, used):
    """名称重复时去重（#2、#3…），Excel 对 sheet 名大小写不敏感，判重按小写。"""
    candidate = name
    i = 2
    while candidate.lower() in used:
        candidate = '%s#%d' % (name[:27], i)
        i += 1
    used.add(candidate.lower())
    return candidate


def px_to_width(px):
    """Univer px → Excel 字符宽（沿用原生成脚本口径：约 1 字符 ≈ 7px）。"""
    return max(5.0, min(255.0, round((px - 5) / 7)))


def write_sheet(ws, sheet):
    cells = sheet.get('cells') or []
    max_row = 0
    max_col = 0
    for cell in cells:
        r = int(cell['r']) + 1
        c = int(cell['c']) + 1
        max_row = max(max_row, r)
        max_col = max(max_col, c)
        formula = cell.get('formula')
        target = ws.cell(row=r, column=c)
        if formula:
            f = str(formula)
            target.value = f if f.startswith('=') else '=' + f
        else:
            target.value = cell.get('v')

    # 合并单元格：起点 (r,c)，跨 rs 行 / cs 列
    for merge in sheet.get('merges') or []:
        r = int(merge['r']) + 1
        c = int(merge['c']) + 1
        rs = max(1, int(merge.get('rs', 1)))
        cs = max(1, int(merge.get('cs', 1)))
        ws.merge_cells(start_row=r, start_column=c,
                       end_row=r + rs - 1, end_column=c + cs - 1)

    # 列宽：colWidths 按 dump 列号存 Univer px；缺省 85px
    col_widths = sheet.get('colWidths') or {}
    for ci in range(1, max_col + 1):
        raw = col_widths.get(str(ci - 1))
        px = float(raw) if isinstance(raw, (int, float)) and raw > 0 else 85.0
        ws.column_dimensions[get_column_letter(ci)].width = px_to_width(px)

    # 全表细边框：覆盖已用区域每个格子（含合并区域内部格）
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).border = ALL_BORDER

    # 表头 / 数据格样式
    for cell in cells:
        target = ws.cell(row=int(cell['r']) + 1, column=int(cell['c']) + 1)
        bg = str(cell.get('bg') or '').lower()
        if bg in HEADER_FILLS:
            rgb = HEADER_FILLS[bg]
            target.fill = PatternFill(fill_type='solid', start_color=rgb, end_color=rgb)
            target.font = HEADER_FONT
            target.alignment = HEADER_ALIGN
        else:
            target.font = DATA_FONT
            target.alignment = DATA_ALIGN


def build_workbook(dump):
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for index, sheet in enumerate(dump.get('sheets') or []):
        raw_name = sheet.get('name') or ''
        name = unique_sheet_name(
            sanitize_sheet_name(raw_name, 'Sheet%d' % (index + 1)), used)
        ws = wb.create_sheet(title=name)
        write_sheet(ws, sheet)
    return wb


def main():
    if not os.path.isdir(DUMP_DIR):
        print('FAIL .dump 目录不存在：%s' % DUMP_DIR)
        sys.exit(1)
    dump_files = sorted(f for f in os.listdir(DUMP_DIR) if f.lower().endswith('.json'))
    if not dump_files:
        print('FAIL .dump 目录为空：%s' % DUMP_DIR)
        sys.exit(1)

    written = 0
    failures = 0
    for fname in dump_files:
        code = fname[:-5]
        try:
            with open(os.path.join(DUMP_DIR, fname), 'r', encoding='utf-8') as fh:
                dump = json.load(fh)
            wb = build_workbook(dump)
            out_path = os.path.join(BASE_DIR, '%s.xlsx' % code)
            wb.save(out_path)
            written += 1
            print('OK   %s: %d sheet -> %s' % (code, len(dump.get('sheets') or []), out_path))
        except Exception as exc:  # noqa: BLE001
            failures += 1
            print('FAIL %s: %s' % (code, exc))

    print('')
    print('════════ xlsx 生成完成 ════════')
    print('合计: %d/%d 个文件, 失败 %d' % (written, len(dump_files), failures))
    sys.exit(0 if failures == 0 and written == len(dump_files) else 1)


if __name__ == '__main__':
    main()
